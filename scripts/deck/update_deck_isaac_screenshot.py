"""Replace the slide-5 Isaac placeholder in the Track A annual progress deck
with the real Isaac Sim trajectory replay screenshot.

Run after scripts/gnm/isaac_live_trajectory_demo.py has produced
assets/deck/isaac_sim_tracka_trajectory_replay.png:

    python3 scripts/deck/update_deck_isaac_screenshot.py
"""

import copy
import shutil
import sys
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu
from pptx.enum.text import PP_ALIGN

REPO = Path(__file__).resolve().parents[2]
IMAGE = REPO / "assets/deck/isaac_sim_tracka_trajectory_replay.png"
DECK = Path("/home/favl/Documents/GNM-VLNVerse_Track_A_Annual_Progress-presentation-deck(new).pptx")
LABEL = "Isaac Sim trajectory replay for Track A image-goal navigation"

SLIDE_INDEX = 4  # slide 5
PLACEHOLDER_TEXT = "Placeholder for Isaac Sim"
INSERT_TEXT = "Insert final trajectory or simulator screenshot here"
PANEL_NAME = "Shape 40"


def remove(shape):
    shape._element.getparent().remove(shape._element)


def main():
    if not IMAGE.exists():
        sys.exit(f"Screenshot not found: {IMAGE}\nRun the Isaac replay demo first.")
    if not DECK.exists():
        sys.exit(f"Deck not found: {DECK}")

    backup = DECK.with_name(DECK.stem + f".bak-{datetime.now():%Y%m%d-%H%M}" + DECK.suffix)
    shutil.copy2(DECK, backup)

    prs = Presentation(DECK)
    slide = prs.slides[SLIDE_INDEX]

    panel = caption = insert_box = None
    for shape in list(slide.shapes):
        text = shape.text_frame.text if shape.has_text_frame else ""
        if PLACEHOLDER_TEXT in text:
            caption = shape
        elif INSERT_TEXT in text:
            insert_box = shape
        elif shape.name == PANEL_NAME:
            panel = shape

    if panel is None or caption is None:
        sys.exit("Slide 5 placeholder shapes not found; deck layout may have changed.")

    # Decorative icon sitting inside the placeholder panel
    for shape in list(slide.shapes):
        if shape.shape_type == 13 and shape is not panel:
            cx = shape.left + shape.width // 2
            cy = shape.top + shape.height // 2
            if (panel.left < cx < panel.left + panel.width
                    and panel.top < cy < panel.top + panel.height):
                remove(shape)

    if insert_box is not None:
        remove(insert_box)

    # Fit the 16:9 screenshot inside the panel, leaving room for the caption
    pad = Emu(60000)
    caption_h = Emu(300000)
    avail_w = panel.width - 2 * pad
    avail_h = panel.height - 2 * pad - caption_h
    img_w = avail_w
    img_h = int(img_w * 9 / 16)
    if img_h > avail_h:
        img_h = avail_h
        img_w = int(img_h * 16 / 9)
    img_left = panel.left + (panel.width - img_w) // 2
    img_top = panel.top + pad

    slide.shapes.add_picture(str(IMAGE), img_left, img_top, img_w, img_h)

    # Reuse the placeholder text box as the caption to keep the deck's fonts/colours
    tf = caption.text_frame
    first = tf.paragraphs[0]
    for para in list(tf.paragraphs[1:]):
        para._p.getparent().remove(para._p)
    for run in list(first.runs[1:]):
        run._r.getparent().remove(run._r)
    if first.runs:
        first.runs[0].text = LABEL
    else:
        first.text = LABEL
    first.alignment = PP_ALIGN.CENTER
    caption.left = panel.left + pad
    caption.top = img_top + img_h + Emu(30000)
    caption.width = panel.width - 2 * pad
    caption.height = caption_h

    prs.save(DECK)
    print(f"Updated slide 5 of {DECK.name}")
    print(f"Backup saved: {backup}")


if __name__ == "__main__":
    main()
